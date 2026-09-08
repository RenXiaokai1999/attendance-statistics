from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from scripts.models import RecordType
from scripts.parse_punch_records import parse_punch_workbook


STATUS_COLORS = {
    RecordType.OFFICIAL_OUT: "FF002060",
    RecordType.COMP_LEAVE: "FF92D050",
    RecordType.LEAVE: "FF000000",
}
MISSING_COLOR = "FFFF0000"
NO_FILL = PatternFill()


@dataclass(frozen=True)
class AnnotationResult:
    output_path: Path
    colored_cells: int
    cleared_cells: int
    unknown_people: set[str]


def _solid(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def annotate_punch_workbook(
    source_path: Path,
    output_path: Path,
    year: int,
    month: int,
    *,
    statuses: dict[tuple[str, date], RecordType] | None = None,
    duty_dates: set[tuple[str, date]] | None = None,
    conflicts: set[tuple[str, date]] | None = None,
    overwrite: bool = False,
) -> AnnotationResult:
    source_path = Path(source_path)
    output_path = Path(output_path)
    if source_path.resolve() == output_path.resolve():
        raise ValueError("标注文件必须另存，不能覆盖原始签卡记录")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"输出文件已存在：{output_path}")

    parsed = parse_punch_workbook(source_path, year, month)
    statuses = statuses or {}
    duty_dates = duty_dates or set()
    conflicts = conflicts or set()
    known_names = set(parsed.name_rows)
    unknown_people = {name for name, _day in statuses if name not in known_names}

    workbook = load_workbook(source_path)
    sheet = workbook.active
    colored_cells = 0
    cleared_cells = 0
    for name, row in parsed.name_rows.items():
        punched = parsed.punched_dates.get(name, set())
        for day, column in parsed.date_columns.items():
            key = (name, day)
            cell = sheet.cell(row, column)
            if day in punched or key in duty_dates or key in conflicts:
                cell.fill = NO_FILL
                cleared_cells += 1
                continue
            kind = statuses.get(key)
            color = STATUS_COLORS.get(kind, MISSING_COLOR)
            cell.fill = _solid(color)
            colored_cells += 1
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return AnnotationResult(
        output_path=output_path,
        colored_cells=colored_cells,
        cleared_cells=cleared_cells,
        unknown_people=unknown_people,
    )
