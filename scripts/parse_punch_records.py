from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from scripts.normalize import normalize_name


DATE_HEADER = re.compile(r"(?<!\d)(\d{1,2})\s*[-/.月]\s*(\d{1,2})(?:日)?(?!\d)")


@dataclass(frozen=True)
class PunchParseResult:
    attendance_days: dict[str, int]
    punched_dates: dict[str, set[date]]
    date_columns: dict[date, int]
    name_rows: dict[str, int]
    header_row: int


def _find_header_row(sheet) -> tuple[int, int]:
    for row in range(1, min(sheet.max_row, 15) + 1):
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row, column).value
            if normalize_name(str(value or "")) == "姓名":
                return row, column
    raise ValueError("签卡记录中未找到“姓名”表头")


def parse_punch_workbook(path: Path, year: int, month: int) -> PunchParseResult:
    workbook = load_workbook(Path(path), data_only=False, read_only=True)
    try:
        sheet = workbook.active
        header_row, name_column = _find_header_row(sheet)
        date_columns: dict[date, int] = {}
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(header_row, column).value
            match = DATE_HEADER.search(str(value or ""))
            if not match:
                continue
            header_month, header_day = map(int, match.groups())
            if header_month != month:
                continue
            try:
                parsed_day = date(year, header_month, header_day)
            except ValueError as exc:
                raise ValueError(f"签卡记录包含无效日期表头：{value}") from exc
            date_columns[parsed_day] = column
        if not date_columns:
            raise ValueError(f"签卡记录中未找到{year}年{month}月日期列")

        punched_dates: dict[str, set[date]] = defaultdict(set)
        name_rows: dict[str, int] = {}
        for row in range(header_row + 1, sheet.max_row + 1):
            name = normalize_name(str(sheet.cell(row, name_column).value or ""))
            if not name:
                continue
            name_rows.setdefault(name, row)
            for day, column in date_columns.items():
                value = sheet.cell(row, column).value
                if value is not None and str(value).strip():
                    punched_dates[name].add(day)
            punched_dates.setdefault(name, set())
        return PunchParseResult(
            attendance_days={name: len(days) for name, days in punched_dates.items()},
            punched_dates=dict(punched_dates),
            date_columns=date_columns,
            name_rows=name_rows,
            header_row=header_row,
        )
    finally:
        workbook.close()

