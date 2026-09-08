from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from scripts.write_attendance_doc import (
    AttendanceSummary, OfficialOutDetail, read_attendance_doc,
    read_official_out_details,
)


def validate_word_output(
    path: Path, expected: dict[str, AttendanceSummary],
    expected_details: list[OfficialOutDetail] | None = None,
) -> None:
    rows = read_attendance_doc(path)
    field_map = {
        "attendance": "出勤",
        "official_out": "因公外出",
        "leave": "请假",
        "all_day_duty": "全日班",
        "day_duty": "白班",
        "comp_leave": "补休",
    }
    for name, summary in expected.items():
        if name not in rows:
            raise ValueError(f"Word输出缺少人员：{name}")
        for field, label in field_map.items():
            expected_value = getattr(summary, field)
            actual = rows[name][label]
            if actual != (str(expected_value) if expected_value else ""):
                raise ValueError(f"Word输出校验失败：{name} {label}")
    if expected_details is not None:
        actual_details = read_official_out_details(path)
        if len(actual_details) != len(expected_details):
            raise ValueError(
                f"Word因公外出明细数量不符：期望{len(expected_details)}，实际{len(actual_details)}"
            )


def validate_excel_output(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        if not workbook.sheetnames or workbook.active.max_row < 2:
            raise ValueError("Excel输出工作表为空")
    finally:
        workbook.close()
