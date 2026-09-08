import hashlib
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from scripts.annotate_punch_workbook import annotate_punch_workbook
from scripts.models import RecordType


class AnnotatePunchWorkbookTests(unittest.TestCase):
    def test_applies_priority_and_colors_without_overwriting_source(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "签卡记录2026-08-01到2026-08-05.xlsx"
            output = root / "签卡记录2026-08-01到2026-08-05_考勤标注.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["考勤序号", "姓名", "工号", "部门", "8-1", "8-2", "8-3", "8-4", "8-5"])
            sheet.append([1, "张子俊", 1, "大同台", "08:00", "", "", "", ""])
            sheet.append([2, "韩胜", 2, "大同台", "", "", "", "", ""])
            yellow = PatternFill("solid", fgColor="FFFF00")
            sheet["E2"].fill = yellow
            sheet["F3"].fill = yellow
            workbook.save(source)
            source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

            annotate_punch_workbook(
                source,
                output,
                2026,
                8,
                statuses={
                    ("张子俊", date(2026, 8, 2)): RecordType.OFFICIAL_OUT,
                    ("张子俊", date(2026, 8, 3)): RecordType.COMP_LEAVE,
                    ("张子俊", date(2026, 8, 4)): RecordType.LEAVE,
                    ("韩胜", date(2026, 8, 2)): RecordType.OFFICIAL_OUT,
                },
                duty_dates={("韩胜", date(2026, 8, 2))},
                conflicts={("韩胜", date(2026, 8, 3))},
            )

            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), source_hash)
            marked = load_workbook(output).active
            self.assertIsNone(marked["E2"].fill.fill_type)
            self.assertEqual(marked["F2"].fill.fgColor.rgb, "FF002060")
            self.assertEqual(marked["G2"].fill.fgColor.rgb, "FF92D050")
            self.assertEqual(marked["H2"].fill.fgColor.rgb, "FF000000")
            self.assertEqual(marked["I2"].fill.fgColor.rgb, "FFFF0000")
            self.assertIsNone(marked["F3"].fill.fill_type)
            self.assertIsNone(marked["G3"].fill.fill_type)
            self.assertEqual(marked["I3"].fill.fgColor.rgb, "FFFF0000")


if __name__ == "__main__":
    unittest.main()
